#!/usr/bin/env python3
"""Parse the client family-members Excel outline into JSON."""

from __future__ import annotations

import json
import re
import sys
from pathlib import Path

from openpyxl import load_workbook

GEN_LABELS = {
    0: "Head of family",
    1: "Generation 1 (child of head)",
    2: "Generation 2 (grandchild)",
    3: "Generation 3 (great-grandchild)",
    4: "Generation 4",
}

FEMALE_OCC = re.compile(r"house\s*wif", re.I)
DIED_OCC = re.compile(r"^died$", re.I)
STATS_B = re.compile(r"(?i)^(members|female|male)\s*-")
COVER_B = re.compile(r"(?i)cover\s*no")
LATE_PREFIX = re.compile(r"(?i)^late\.?\s*")
NUM_PREFIX = re.compile(r"^\d+\.\s*")
DOB_DMY = re.compile(r"^(\d{1,2})\.(\d{1,2})\.(\d{2,4})$")
AGE_YEARS = re.compile(r"(?i)(\d+)\s*years?")


def cell(row, i) -> str:
    if i >= len(row):
        return ""
    v = row[i]
    if v is None:
        return ""
    return str(v).strip()


def blank(value: str | None) -> bool:
    return not value or value in {"-", "–", "—", "-'"}


def clean_name(name: str) -> tuple[str, bool]:
    name = NUM_PREFIX.sub("", name).strip()
    deceased = bool(LATE_PREFIX.match(name))
    name = LATE_PREFIX.sub("", name)
    name = re.sub(r"\s+", " ", name).strip(" .")
    return name, deceased


def parse_dob(raw: str) -> tuple[str | None, str | None]:
    if blank(raw):
        return None, None
    m = DOB_DMY.match(raw)
    if m:
        d, mo, y = m.groups()
        if len(y) == 2:
            y = ("19" + y) if int(y) > 30 else ("20" + y)
        try:
            return f"{int(y):04d}-{int(mo):02d}-{int(d):02d}", None
        except ValueError:
            return None, raw
    m = AGE_YEARS.search(raw)
    if m:
        return None, f"Age {m.group(1)} (as of 2026-09-03)"
    if re.fullmatch(r"\d{1,3}", raw):
        return None, f"Age {raw} (as of 2026-09-03)"
    return None, raw


def infer_gender(occupation: str) -> str | None:
    if FEMALE_OCC.search(occupation or ""):
        return "Female"
    return None


def relationship_for(gen: int, gender: str | None) -> str:
    if gen <= 0:
        return "other"
    if gen == 1:
        if gender == "Female":
            return "daughter"
        if gender == "Male":
            return "son"
        return "other"
    if gen == 2:
        if gender == "Female":
            return "granddaughter"
        if gender == "Male":
            return "grandson"
        return "other"
    if gen == 3:
        if gender == "Female":
            return "great_granddaughter"
        if gender == "Male":
            return "great_grandson"
        return "other"
    return "other"


def remarks_for(gen: int, occupation: str, age_note: str | None) -> str | None:
    parts = [GEN_LABELS.get(gen, f"Generation {gen}")]
    if occupation:
        parts.append(occupation)
    if age_note:
        parts.append(age_note)
    text = " | ".join(parts)
    return text[:1000] if text else None


def parse_workbook(path: Path) -> dict:
    wb = load_workbook(path, data_only=True)
    ws = wb["Sheet1"]
    families: list[dict] = []
    current: dict | None = None

    def start_cover(sl: str, cover: str):
        nonlocal current
        m = re.search(r"(\d+)", cover)
        number = m.group(1).zfill(2) if m else sl.zfill(2)
        current = {
            "sl": sl,
            "cover": cover,
            "cover_no": number,
            "family_no": f"COVER-{number}",
            "family_name": "",
            "username": f"cover{number}",
            "members": [],
        }

    def add_member(**kwargs):
        assert current is not None
        name = kwargs["full_name"]
        if not name or len(name) < 2:
            return
        current["members"].append(kwargs)

    def patch_head(**kwargs):
        assert current is not None
        for member in current["members"]:
            if member.get("is_head"):
                for key, value in kwargs.items():
                    if value not in (None, "", False) or key == "is_deceased":
                        if key == "is_deceased":
                            member[key] = member.get("is_deceased") or value
                        elif not member.get(key):
                            member[key] = value
                return

    for row in ws.iter_rows(min_row=3, max_col=10, values_only=True):
        vals = [cell(row, i) for i in range(10)]
        if not any(vals):
            continue
        b = vals[1]
        if COVER_B.search(b or ""):
            if current:
                families.append(current)
            start_cover(vals[0], b)
            continue
        if current is None:
            continue

        dob, age_note = parse_dob(vals[6])
        occupation = "" if blank(vals[7]) else vals[7]
        phone = None if blank(vals[8]) else vals[8]
        deceased_occ = bool(DIED_OCC.match(occupation))
        gender = infer_gender(occupation)

        gen_names: list[tuple[int, str, bool]] = []
        for gen, idx in enumerate([2, 3, 4, 5], start=1):
            raw = vals[idx]
            if blank(raw):
                continue
            name, deceased = clean_name(raw)
            if name:
                gen_names.append((gen, name, deceased))

        if b and not STATS_B.match(b) and not COVER_B.search(b):
            head_name, head_deceased = clean_name(b)
            if head_name and not current["family_name"]:
                current["family_name"] = head_name
                add_member(
                    full_name=head_name,
                    gen=0,
                    is_head=True,
                    is_deceased=head_deceased or deceased_occ,
                    relationship="other",
                    gender=gender if not gen_names else None,
                    date_of_birth=dob if not gen_names else None,
                    contact=phone if not gen_names else None,
                    remarks=remarks_for(0, occupation if not gen_names else "", age_note if not gen_names else None),
                )
            elif head_name and head_name == current["family_name"] and not gen_names:
                patch_head(
                    is_deceased=head_deceased or deceased_occ,
                    gender=gender,
                    date_of_birth=dob,
                    contact=phone,
                    remarks=remarks_for(0, occupation, age_note),
                )

        if not gen_names:
            continue

        deepest = max(g for g, _, _ in gen_names)
        for gen, name, deceased in gen_names:
            is_deepest = gen == deepest
            g = gender if is_deepest else None
            add_member(
                full_name=name,
                gen=gen,
                is_head=False,
                is_deceased=deceased or (deceased_occ if is_deepest else False),
                relationship=relationship_for(gen, g),
                gender=g,
                date_of_birth=dob if is_deepest else None,
                contact=phone if is_deepest else None,
                remarks=remarks_for(gen, occupation if is_deepest else "", age_note if is_deepest else None),
            )

    if current:
        families.append(current)

    for family in families:
        if not family["family_name"]:
            family["family_name"] = family["family_no"]
        phones = [m.get("contact") for m in family["members"] if m.get("contact")]
        family["contact_phone"] = phones[0] if phones else None

    return {
        "source": str(path),
        "title": "PANAMBUR ANDOTA ASHTA DAIVAGALA SEVA TRUST (SALIAN KUTUMBA) NADOORI",
        "families": families,
        "family_count": len(families),
        "member_count": sum(len(f["members"]) for f in families),
    }


def main() -> int:
    path = Path(sys.argv[1] if len(sys.argv) > 1 else "")
    if not path.is_file():
        print(f"Excel file not found: {path}", file=sys.stderr)
        return 1
    json.dump(parse_workbook(path), sys.stdout, ensure_ascii=False, indent=2)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
